"""Fixtures compartilhadas: XLSX mínimo no formato SSP."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook


def write_celulares_xlsx(path: Path, rows: list[dict] | None = None) -> Path:
    wb = Workbook()
    wb.active.title = "METODOLOGIA"
    wb.active["A1"] = "Manual de interpretação (fixture)"
    dicio = wb.create_sheet("DICIONARIO DE DADOS")
    dicio["A1"] = "LATITUDE"
    dicio["B1"] = "coordenada"
    sheet = wb.create_sheet("CELULAR_2026")
    columns = [
        "LATITUDE",
        "LONGITUDE",
        "RUBRICA",
        "DATA_OCORRENCIA_BO",
        "HORA_OCORRENCIA",
        "DESCR_PERIODO",
        "BAIRRO",
        "CIDADE",
        "LOGRADOURO",
        "NOME_DELEGACIA",
        "MARCA_OBJETO",
        "NUM_BO",
        "VERSAO",
        "ANO_BO",
    ]
    sheet.append(columns)
    default = [
        {
            "LATITUDE": -23.5505,
            "LONGITUDE": -46.6333,
            "RUBRICA": "Furto (art. 155)",
            "DATA_OCORRENCIA_BO": "2026-01-15",
            "HORA_OCORRENCIA": "10:30:00",
            "DESCR_PERIODO": None,
            "BAIRRO": "CENTRO",
            "CIDADE": "S.PAULO",
            "LOGRADOURO": "Rua A",
            "NOME_DELEGACIA": "01 DP",
            "MARCA_OBJETO": "Apple",
            "NUM_BO": "BO-1",
            "VERSAO": 1,
            "ANO_BO": 2026,
        },
        {
            "LATITUDE": -23.5616,
            "LONGITUDE": -46.6558,
            "RUBRICA": "Roubo (art. 157)",
            "DATA_OCORRENCIA_BO": "2026-02-10",
            "HORA_OCORRENCIA": "22:00:00",
            "DESCR_PERIODO": None,
            "BAIRRO": "PAULISTA",
            "CIDADE": "S.PAULO",
            "LOGRADOURO": "Av Paulista",
            "NOME_DELEGACIA": "78 DP",
            "MARCA_OBJETO": "Samsung",
            "NUM_BO": "BO-2",
            "VERSAO": 1,
            "ANO_BO": 2026,
        },
        {
            "LATITUDE": None,
            "LONGITUDE": None,
            "RUBRICA": "Furto",
            "DATA_OCORRENCIA_BO": "2026-03-01",
            "HORA_OCORRENCIA": None,
            "DESCR_PERIODO": "Tarde",
            "BAIRRO": "X",
            "CIDADE": "S.PAULO",
            "LOGRADOURO": None,
            "NOME_DELEGACIA": "01 DP",
            "MARCA_OBJETO": None,
            "NUM_BO": "BO-3",
            "VERSAO": 1,
            "ANO_BO": 2026,
        },
        {
            "LATITUDE": 0,
            "LONGITUDE": 0,
            "RUBRICA": "Furto",
            "DATA_OCORRENCIA_BO": "2026-03-02",
            "HORA_OCORRENCIA": None,
            "DESCR_PERIODO": None,
            "BAIRRO": "Y",
            "CIDADE": "S.PAULO",
            "LOGRADOURO": None,
            "NOME_DELEGACIA": "01 DP",
            "MARCA_OBJETO": None,
            "NUM_BO": "BO-4",
            "VERSAO": 1,
            "ANO_BO": 2026,
        },
        {
            "LATITUDE": -23.5480,
            "LONGITUDE": -46.6400,
            "RUBRICA": "Estelionato",
            "DATA_OCORRENCIA_BO": "2026-04-05",
            "HORA_OCORRENCIA": "14:00:00",
            "DESCR_PERIODO": None,
            "BAIRRO": "BELA VISTA",
            "CIDADE": "S.PAULO",
            "LOGRADOURO": "Rua B",
            "NOME_DELEGACIA": "04 DP",
            "MARCA_OBJETO": None,
            "NUM_BO": "BO-5",
            "VERSAO": 1,
            "ANO_BO": 2026,
        },
    ]
    for row in rows or default:
        sheet.append([row.get(c) for c in columns])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_spdados_xlsx(path: Path) -> Path:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "DADOS"
    columns = [
        "LATITUDE",
        "LONGITUDE",
        "NATUREZA_APURADA",
        "DATA_OCORRENCIA",
        "HORA_OCORRENCIA",
        "BAIRRO",
        "CIDADE",
        "DELEGACIA_ELABORACAO",
        "NUM_BO",
        "ANO_BO",
    ]
    sheet.append(columns)
    sheet.append([-23.5505, -46.6333, "Furto", "2026-01-20", "09:00:00", "SE", "SAO PAULO", "01 DP", "99", 2026])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


@pytest.fixture
def celulares_xlsx(tmp_path: Path) -> Path:
    return write_celulares_xlsx(tmp_path / "CelularesSubtraidos_2026.xlsx")


@pytest.fixture
def spdados_xlsx(tmp_path: Path) -> Path:
    return write_spdados_xlsx(tmp_path / "SPDadosCriminais_2026.xlsx")
